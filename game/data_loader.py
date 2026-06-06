"""Game data loader: reads data/game_data.xlsx and returns config-compatible dict.

Content data (monsters, items, skills, shop, etc.) is declared in Excel
and compiled into Python objects at load time.  Engine constants stay in
config.py.
"""

from __future__ import annotations

import os
from typing import Any, Callable

# ---------------------------------------------------------------------------
# Helper (mirrors utils._set so the loader is self-contained for effects)
# ---------------------------------------------------------------------------

def _get(obj: object, attr: str, default: Any = None) -> Any:
    """Get attribute value from dict or object."""
    if isinstance(obj, dict):
        return obj.get(attr, default)
    return getattr(obj, attr, default)


def _set(obj: object, attr: str, val: Any) -> None:
    """Set attribute value on dict or object."""
    if isinstance(obj, dict):
        obj[attr] = val
    else:
        setattr(obj, attr, val)


# ---------------------------------------------------------------------------
# Skill effect compiler
# ---------------------------------------------------------------------------

def _build_effect(row: dict[str, Any]) -> Callable:
    """Compile one skill row (from Excel) into a callable *effect* fn.

    The returned lambda has signature ``fn(player_data) -> None``.

    Supported **effect_type** values
    ---------------------------------
    * ``"add_attr"``       – single attribute increment
    * ``"add_attr_pair"``  – two attributes incremented together
    * ``"set_flag"``       – set a boolean flag to True
    * ``"stacked_mult"``   – multiplicative value with stack counter
    """
    etype: str = row["effect_type"]

    if etype == "add_attr":
        attr: str = row["target_attr"]
        val = row["value"]
        default = 0.0 if isinstance(val, float) else 0
        return lambda p, a=attr, v=val, d=default: _set(
            p, a, _get(p, a, d) + v
        )

    elif etype == "add_attr_pair":
        a1, v1 = row["attr1"], row["value1"]
        a2, v2 = row["attr2"], row["value2"]
        return lambda p, _a1=a1, _v1=v1, _a2=a2, _v2=v2: (
            _set(p, _a1, _get(p, _a1, 0) + _v1),
            _set(p, _a2, _get(p, _a2, 0) + _v2),
        )

    elif etype == "set_flag":
        fname: str = row["flag_name"]
        return lambda p, f=fname: _set(p, f, True)

    elif etype == "stacked_mult":
        attr = row["target_attr"]
        base = row["base"]
        sattr = row["stack_attr"]
        return lambda p, a=attr, b=base, s=sattr: _set(
            p, a, b * (1 + _get(p, s, 0))
        )

    else:
        raise ValueError(f"Unknown effect_type: {etype!r}")


# ---------------------------------------------------------------------------
# Sheet readers
# ---------------------------------------------------------------------------

def _read_header(ws) -> list[str]:
    """Return header row values from worksheet."""
    return [cell.value for cell in ws[1]]


def _row_to_dict(ws, row_idx: int, headers: list[str]) -> dict[str, Any]:
    """Convert a worksheet data row to {header: value} dict."""
    values = [cell.value for cell in ws[row_idx]]
    return {
        h: v for h, v in zip(headers, values)
        if h is not None
    }


def _load_constants(ws) -> dict[str, Any]:
    headers = _read_header(ws)
    result: dict[str, Any] = {}
    for r in range(2, ws.max_row + 1):
        row = _row_to_dict(ws, r, headers)
        key = row.get("key")
        if key is None:
            continue
        vt = row.get("value_type", "")
        val = row.get("value")
        if vt == "int":
            result[key] = int(val)
        elif vt == "float":
            result[key] = float(val)
        else:
            result[key] = val
    return result


def _load_tiles(ws) -> dict[str, str]:
    headers = _read_header(ws)
    result: dict[str, str] = {}
    for r in range(2, ws.max_row + 1):
        row = _row_to_dict(ws, r, headers)
        k = row.get("key")
        if k is not None:
            result[k] = str(row.get("char", ""))
    return result


def _load_monsters(ws) -> list[dict]:
    headers = _read_header(ws)
    monsters: list[dict] = []
    for r in range(2, ws.max_row + 1):
        row = _row_to_dict(ws, r, headers)
        if row.get("char"):
            monsters.append({
                "char": row["char"],
                "name": row.get("name", ""),
                "hp": int(row.get("hp", 0)),
                "attack": int(row.get("attack", 0)),
                "defense": int(row.get("defense", 0)),
                "exp": int(row.get("exp", 0)),
                "gold": int(row.get("gold", 0)),
            })
    return monsters


def _load_items(ws) -> list[dict]:
    headers = _read_header(ws)
    items: list[dict] = []
    for r in range(2, ws.max_row + 1):
        row = _row_to_dict(ws, r, headers)
        if row.get("char"):
            items.append({
                "char": row["char"],
                "name": row.get("name", ""),
                "type": row.get("type", ""),
                "value": int(row.get("value", 0)),
            })
    return items


def _load_boss(ws) -> dict:
    headers = _read_header(ws)
    if ws.max_row < 2:
        return {}
    row = _row_to_dict(ws, 2, headers)
    return {
        "char": row.get("char", "B"),
        "name": row.get("name", ""),
        "hp": int(row.get("hp", 0)),
        "attack": int(row.get("attack", 0)),
        "defense": int(row.get("defense", 0)),
        "exp": int(row.get("exp", 0)),
        "gold": int(row.get("gold", 0)),
    }


def _load_player_base(ws) -> dict:
    headers = _read_header(ws)
    if ws.max_row < 2:
        return {"hp": 100, "attack": 10, "defense": 5}
    row = _row_to_dict(ws, 2, headers)
    return {
        "hp": int(row.get("hp", 100)),
        "attack": int(row.get("attack", 10)),
        "defense": int(row.get("defense", 5)),
    }


def _load_shop_items(ws) -> list[dict]:
    headers = _read_header(ws)
    items: list[dict] = []
    for r in range(2, ws.max_row + 1):
        row = _row_to_dict(ws, r, headers)
        if row.get("id") is None:
            continue
        item: dict[str, Any] = {
            "id": int(row["id"]),
            "name": row.get("name", ""),
            "desc": row.get("desc", ""),
            "price": int(row.get("price", 0)),
            "stock": int(row.get("stock", 0)),
            "type": row.get("type", ""),
        }
        # Type-specific fields — kept as-is; world.py buy_item() uses id/type dispatch
        item["effect_target"] = row.get("effect_target") or ""
        item["effect_value"] = _to_float(row.get("effect_value"))
        item["effect_duration"] = _to_int(row.get("effect_duration"))
        item["permanent_attr"] = row.get("permanent_attr") or ""
        item["permanent_value"] = _to_int(row.get("permanent_value"))
        item["heal_value"] = _to_int(row.get("heal_value"))
        item["key_value"] = _to_int(row.get("key_value"))
        items.append(item)
    return items


def _load_skills(ws) -> list[dict]:
    headers = _read_header(ws)
    skills: list[dict] = []
    for r in range(2, ws.max_row + 1):
        row = _row_to_dict(ws, r, headers)
        if row.get("id") is None:
            continue
        skill_def: dict[str, Any] = {
            "id": int(row["id"]),
            "name": row.get("name", ""),
            "desc": row.get("desc", ""),
            "max_level": int(row.get("max_level", 1)),
            "effect_type": row.get("effect_type", ""),
        }
        # Parameters for compiler
        skill_def["target_attr"] = row.get("target_attr") or ""
        skill_def["value"] = _to_float(row.get("value"))
        skill_def["attr1"] = row.get("attr1") or ""
        skill_def["value1"] = _to_float(row.get("value1"))
        skill_def["attr2"] = row.get("attr2") or ""
        skill_def["value2"] = _to_float(row.get("value2"))
        skill_def["flag_name"] = row.get("flag_name") or ""
        skill_def["base"] = _to_float(row.get("base"))
        skill_def["stack_attr"] = row.get("stack_attr") or ""
        # Compile the effect lambda
        skill_def["effect"] = _build_effect(skill_def)
        skills.append(skill_def)
    return skills


# ---------------------------------------------------------------------------
# Tiny helpers
# ---------------------------------------------------------------------------

def _to_float(val, default=None):
    if val is None:
        return default
    try:
        return float(val)
    except (TypeError, ValueError):
        return default


def _to_int(val, default=None):
    if val is None:
        return default
    try:
        return int(float(val))
    except (TypeError, ValueError):
        return default


# ---------------------------------------------------------------------------
# Main loader class
# ---------------------------------------------------------------------------

class GameDataLoader:
    """Read *data/game_data.xlsx* and produce a config-compatible dict."""

    def __init__(self, xlsx_path: str | None = None):
        if xlsx_path is None:
            # data_loader.py lives in game/, project root is parent of game/
            _game_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            xlsx_path = os.path.join(_game_dir, "data", "game_data.xlsx")
        self.xlsx_path: str = xlsx_path

    # ---- public API ----

    def load_all(self) -> dict[str, Any]:
        """Load all sheets and return a merged config-compatible dictionary.

        Keys match the UPPERCASE variable names used by config.py / world.py:
        MONSTER_TYPES, BOSS_TEMPLATE, ITEM_TYPES, PLAYER_BASE,
        SHOP_ITEMS, ALL_SKILLS, plus engine constants and TILES.
        """
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ImportError(
                "openpyxl is required.  Install it with: pip install openpyxl"
            )

        if not os.path.isfile(self.xlsx_path):
            raise FileNotFoundError(
                f"Game data file not found: {self.xlsx_path}\n"
                "Please ensure data/game_data.xlsx exists in the project root."
            )

        wb = load_workbook(self.xlsx_path, read_only=True, data_only=True)

        result: dict[str, Any] = {}

        # Engine constants
        if "constants" in wb.sheetnames:
            result.update(_load_constants(wb["constants"]))

        # Tile definitions
        if "tiles" in wb.sheetnames:
            result["TILES"] = _load_tiles(wb["tiles"])

        # Content data (uppercase keys to match config.py convention)
        if "monsters" in wb.sheetnames:
            result["MONSTER_TYPES"] = _load_monsters(wb["monsters"])

        if "items" in wb.sheetnames:
            result["ITEM_TYPES"] = _load_items(wb["items"])

        if "boss" in wb.sheetnames:
            result["BOSS_TEMPLATE"] = _load_boss(wb["boss"])

        if "player_base" in wb.sheetnames:
            result["PLAYER_BASE"] = _load_player_base(wb["player_base"])

        if "shop_items" in wb.sheetnames:
            result["SHOP_ITEMS"] = _load_shop_items(wb["shop_items"])

        if "skills" in wb.sheetnames:
            result["ALL_SKILLS"] = _load_skills(wb["skills"])

        wb.close()
        return result
